"""
Script to generate the comprehensive publication-grade Dataset Description Table Excel workbook.
Supports the paper:
'Algorithm-Driven Alumni Management: An Intelligent Approach'
Central Framework: Integrated Alumni Intelligence Framework (IAIF)
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # -------------------------------------------------------------
    # Styling Palette & Definitions
    # -------------------------------------------------------------
    FONT_FAMILY = "Segoe UI"
    
    # Colors
    NAVY_DARK = "1B365D"      # Title / Primary Header Fill
    STEEL_BLUE = "2B547E"     # Subheader / Section Fill
    ACCENT_BLUE = "DCE6F1"    # Accent / Highlight Box Fill
    ACCENT_LIGHT = "F1F5F9"   # Alternating Row Tint
    WHITE = "FFFFFF"
    DARK_TEXT = "1E293B"
    MUTED_TEXT = "475569"
    BORDER_COLOR = "CBD5E1"
    
    # Fonts
    title_font = Font(name=FONT_FAMILY, size=14, bold=True, color=WHITE)
    subtitle_font = Font(name=FONT_FAMILY, size=10, italic=True, color="E2E8F0")
    section_font = Font(name=FONT_FAMILY, size=11, bold=True, color=WHITE)
    header_font = Font(name=FONT_FAMILY, size=10, bold=True, color=WHITE)
    data_font = Font(name=FONT_FAMILY, size=9.5, color=DARK_TEXT)
    data_bold_font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=DARK_TEXT)
    kpi_val_font = Font(name=FONT_FAMILY, size=14, bold=True, color=NAVY_DARK)
    kpi_lbl_font = Font(name=FONT_FAMILY, size=9, bold=True, color=MUTED_TEXT)
    total_font = Font(name=FONT_FAMILY, size=10, bold=True, color=DARK_TEXT)
    
    # Fills
    title_fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    section_fill = PatternFill(start_color=STEEL_BLUE, end_color=STEEL_BLUE, fill_type="solid")
    header_fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    header_sub_fill = PatternFill(start_color="334E68", end_color="334E68", fill_type="solid")
    alt_fill = PatternFill(start_color=ACCENT_LIGHT, end_color=ACCENT_LIGHT, fill_type="solid")
    white_fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
    kpi_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    total_fill = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
    
    # Borders
    thin_border_side = Side(border_style="thin", color=BORDER_COLOR)
    double_bottom_side = Side(border_style="double", color=BORDER_COLOR)
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    kpi_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    total_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_bottom_side)
    
    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_kpi_val = Alignment(horizontal="center", vertical="center")
    align_kpi_lbl = Alignment(horizontal="center", vertical="center")

    def apply_title_banner(ws, title_text, subtitle_text, num_cols):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        c1 = ws.cell(row=1, column=1, value=title_text)
        c1.font = title_font
        c1.fill = title_fill
        c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28
        
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
        c2 = ws.cell(row=2, column=1, value=subtitle_text)
        c2.font = subtitle_font
        c2.fill = title_fill
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 20

    def style_headers(ws, row_idx, headers, fill=header_fill):
        ws.row_dimensions[row_idx].height = 26
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = fill
            cell.alignment = align_header
            cell.border = cell_border

    def auto_fit_columns(ws, min_widths=None, max_width=55):
        if min_widths is None:
            min_widths = {}
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                # ignore merged title banner in length check
                if cell.row in [1, 2]:
                    continue
                val = str(cell.value or '')
                if '\n' in val:
                    val = max(val.split('\n'), key=len)
                max_len = max(max_len, len(val))
            calc_width = max(max_len + 4, min_widths.get(col_letter, 12))
            ws.column_dimensions[col_letter].width = min(calc_width, max_width)

    # =========================================================================
    # SHEET 1: Dataset_Catalog_Summary
    # =========================================================================
    ws1 = wb.create_sheet(title="Dataset_Catalog_Summary")
    ws1.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws1,
        "INTEGRATED ALUMNI INTELLIGENCE FRAMEWORK (IAIF) - DATASET CATALOG & SYSTEM OVERVIEW",
        "Algorithm-Driven Alumni Management: Architectural Schema, Scale Metrics, and Relational Taxonomy",
        10
    )
    
    # KPI Summary Cards (Row 4 & 5)
    kpis = [
        ("Total Alumni Profiles", "6,000", 1, 2),
        ("Total Interactions", "107,487", 3, 4),
        ("Network Connections", "29,984", 5, 6),
        ("Events Cataloged", "150", 7, 8),
        ("Certified Mentors", "600", 9, 10)
    ]
    
    ws1.row_dimensions[4].height = 24
    ws1.row_dimensions[5].height = 18
    
    for lbl, val, c_start, c_end in kpis:
        ws1.merge_cells(start_row=4, start_column=c_start, end_row=4, end_column=c_end)
        vc = ws1.cell(row=4, column=c_start, value=val)
        vc.font = kpi_val_font
        vc.fill = kpi_fill
        vc.alignment = align_kpi_val
        
        ws1.merge_cells(start_row=5, start_column=c_start, end_row=5, end_column=c_end)
        lc = ws1.cell(row=5, column=c_start, value=lbl)
        lc.font = kpi_lbl_font
        lc.fill = kpi_fill
        lc.alignment = align_kpi_lbl
        
        for r in [4, 5]:
            for c in range(c_start, c_end + 1):
                ws1.cell(row=r, column=c).border = kpi_border
                ws1.cell(row=r, column=c).fill = kpi_fill
    
    # Master Dataset Summary Table Header
    headers1 = [
        "Table ID", "Entity / Dataset Name", "File Name", "Total Records (Rows)", 
        "Attributes (Cols)", "Primary Key", "Foreign Key(s)", "Temporal Scope", 
        "Data Sparsity / Missing %", "Primary Role in Central IAIF Architecture"
    ]
    style_headers(ws1, 7, headers1)
    
    data1 = [
        ("DS-01", "Alumni Profile", "Alumni_Profile.csv", 6000, 10, "alumni_id", "None (Master Entity)", "Static / Longitudinal (2005-2023)", "0.0% (Complete)", "Demographic & educational foundations for Multi-View Representation Learning (MVAR-Net) and Behavioral Segmentation (BAAC)."),
        ("DS-02", "Alumni Interaction", "Alumni_Interaction.csv", 107487, 4, "Composite (alumni_id, month, action)", "alumni_id -> Alumni_Profile", "24 Months (18 Obs / 6 Holdout)", "0.0% (Complete)", "Dynamic temporal activity tracking, longitudinal sequence modeling, and engagement/disengagement state forecasting (EEN)."),
        ("DS-03", "Institutional Events", "Event.csv", 150, 6, "event_id", "None (Target Entity)", "24 Months Continuous Schedule", "0.0% (Complete)", "Candidate catalog for personalized multi-task event recommendations and skill-gap matching (RAOD + ACRN)."),
        ("DS-04", "Alumni Mentors", "Mentor.csv", 600, 6, "mentor_id", "mentor_id -> Alumni_Profile", "Active Capacity Pool", "0.0% (Complete)", "Senior professional pool for mentorship pairing, career alignment matching, and capacity-constrained dispatching."),
        ("DS-05", "Alumni Network Graph", "Alumni_Network.csv", 29984, 3, "Composite (source, target)", "source/target -> Alumni_Profile", "Weighted Relational Graph", "0.0% (Complete)", "Peer connectivity topology for Dynamic Relationship Modeling (DRMN) and affinity-driven networking recommendations.")
    ]
    
    cur_row = 8
    for row_data in data1:
        ws1.row_dimensions[cur_row].height = 30
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx in [1, 2, 4] else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [1, 6]:
                cell.alignment = align_center
            elif col_idx in [4, 5]:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            elif col_idx in [9]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
        cur_row += 1
        
    # Total row
    ws1.row_dimensions[cur_row].height = 24
    totals = ["Total", "5 Relational Tables", "All Files in /data", 144221, 29, "-", "-", "24-Month Temporal Horizon", "0.0% Missing", "Full IAIF End-to-End Pipeline Support"]
    for col_idx, val in enumerate(totals, 1):
        cell = ws1.cell(row=cur_row, column=col_idx, value=val)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = total_border
        if col_idx in [4, 5]:
            cell.alignment = align_right
            cell.number_format = "#,##0"
        elif col_idx in [1, 6, 7, 9]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left
            
    # System Architecture Integration Notes Section
    cur_row += 2
    ws1.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=10)
    sec_cell = ws1.cell(row=cur_row, column=1, value="IAIF WORKFLOW STAGES AND DATASET CONSUMPTION MAPPING")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[cur_row].height = 24
    
    workflow_steps = [
        ("Step 1: Synthetic Generation", "Alumni_Profile.csv, Alumni_Interaction.csv, Event.csv, Mentor.csv, Alumni_Network.csv", "Generates realistic dependency-aware synthetic distributions adhering to archetypes, career paths, and 24-month activity dynamics."),
        ("Step 2: Fusion & Preprocessing", "All 5 Datasets", "Integrates multi-table records, extracts 18-month longitudinal aggregates, prevents temporal leakage, and partitions train/val/test splits (70/15/15)."),
        ("Step 3: Multi-View Learning (MVAR-Net)", "Profile + Career/Skills + Behavioral Views", "Encodes distinct information modalities into 64-dimensional unified latent vector representations with attention-based view fusion."),
        ("Step 4: Behavioral Segmentation (BAAC)", "Unified Latents + Career/Behavioral Vectors", "Adaptive weighted K-Means clustering groups alumni into 4 distinct segments (Silhouette=0.1901, Calinski-Harabasz=1634.33)."),
        ("Step 5: Engagement Prediction (EEN)", "Aggregated 18-Month Features + Holdout 19-24 Labels", "Predicts multi-tier engagement state (Low/Med/High) and binary disengagement risk over future 6 months (Accuracy=88.33%, ROC-AUC=0.8409)."),
        ("Step 6: Multitask Recommender (DRMN+ACRN)", "Event.csv, Mentor.csv, Alumni_Network.csv + Master Features", "Generates and contextually ranks personalized event recommendations (P@5=0.979), mentor matches (P@5=0.960), and peer networking connections (P@5=0.940)."),
        ("Step 7: Decision Support (AAIE)", "EEN Predictions + BAAC Clusters + ACRN Rankings", "Converts model inference into actionable operational prioritization queues across 5 institutional intervention tiers.")
    ]
    
    cur_row += 1
    style_headers(ws1, cur_row, ["Workflow Step", "Datasets Consumed", "Functional Objectives and Methodological Output"], fill=header_sub_fill)
    ws1.merge_cells(start_row=cur_row, start_column=3, end_row=cur_row, end_column=10)
    
    for step, dsets, desc in workflow_steps:
        cur_row += 1
        ws1.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        
        c_step = ws1.cell(row=cur_row, column=1, value=step)
        c_step.font = data_bold_font
        c_step.fill = fill
        c_step.border = cell_border
        c_step.alignment = align_left
        
        c_ds = ws1.cell(row=cur_row, column=2, value=dsets)
        c_ds.font = data_font
        c_ds.fill = fill
        c_ds.border = cell_border
        c_ds.alignment = align_left
        
        ws1.merge_cells(start_row=cur_row, start_column=3, end_row=cur_row, end_column=10)
        c_desc = ws1.cell(row=cur_row, column=3, value=desc)
        c_desc.font = data_font
        c_desc.fill = fill
        c_desc.alignment = align_left
        for c_idx in range(3, 11):
            ws1.cell(row=cur_row, column=c_idx).border = cell_border
            ws1.cell(row=cur_row, column=c_idx).fill = fill

    auto_fit_columns(ws1, min_widths={"A": 12, "B": 24, "C": 22, "D": 22, "E": 18, "F": 18, "G": 22, "H": 25, "I": 20, "J": 45}, max_width=50)

    # =========================================================================
    # SHEET 2: Master_Feature_Dictionary
    # =========================================================================
    ws2 = wb.create_sheet(title="Master_Feature_Dictionary")
    ws2.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws2,
        "MASTER FEATURE DICTIONARY - COMPLETE MULTI-TABLE ATTRIBUTE REPOSITORY",
        "Exhaustive Specification of Raw and Relational Features, Data Types, Null Profiles, Permitted Domains, and Pipeline Roles",
        10
    )
    
    headers2 = [
        "Table Name", "Attribute Name", "Physical Data Type", "Measurement Scale", 
        "Null Count (%)", "Distinct Values", "Permitted Domain / Allowed Values", 
        "Operational Description & Semantics", "Sample Values", "IAIF Module & Feature View Role"
    ]
    style_headers(ws2, 4, headers2)
    
    master_features = [
        # Alumni_Profile.csv
        ("Alumni_Profile.csv", "alumni_id", "String (Object)", "Nominal (Identifier)", "0 (0.0%)", 6000, "ALU_00001 to ALU_06000", "Unique primary key designating each individual alumni profile.", "ALU_00001, ALU_00002", "Primary Entity Key (Relational Anchor)"),
        ("Alumni_Profile.csv", "archetype", "Integer", "Nominal (Category)", "0 (0.0%)", 4, "{0, 1, 2, 3}", "Behavioral persona segment: 0=Early Tech, 1=Mid Specialist, 2=Senior Exec, 3=Dormant/At-Risk.", "0, 1, 2, 3", "Synthetic Generating Ground-Truth / Persona Baseline"),
        ("Alumni_Profile.csv", "graduation_year", "Integer", "Interval", "0 (0.0%)", 19, "[2005, 2023]", "Calendar year of degree completion / institutional graduation.", "2009, 2014, 2022", "MVAR-Net: Profile View (Demographic Encoding)"),
        ("Alumni_Profile.csv", "degree", "String (Object)", "Nominal", "0 (0.0%)", 6, "{B.Tech, M.Tech, B.Sc, M.Sc, MBA, Ph.D}", "Highest completed academic credential awarded by the institution.", "B.Tech, M.Tech, MBA", "MVAR-Net: Profile View (One-Hot Encoded)"),
        ("Alumni_Profile.csv", "major", "String (Object)", "Nominal", "0 (0.0%)", 6, "{Computer Science, Data Science, Electrical Eng, Mechanical Eng, Business Admin, Bioengineering}", "Primary academic discipline and major field of undergraduate/graduate study.", "Computer Science, Data Science", "MVAR-Net: Profile View / DRMN Relational Match"),
        ("Alumni_Profile.csv", "industry", "String (Object)", "Nominal", "0 (0.0%)", 6, "{Technology, Finance, Healthcare, Automotive, Consulting, Education}", "Current employment industrial vertical and economic sector.", "Technology, Consulting, Finance", "MVAR-Net: Career View / DRMN Industry Match"),
        ("Alumni_Profile.csv", "years_experience", "Integer", "Ratio", "0 (0.0%)", 19, "[1, 19] Years", "Total cumulative professional work experience post-graduation.", "5, 10, 15", "MVAR-Net: Profile & Career View (Min-Max Scaled)"),
        ("Alumni_Profile.csv", "seniority", "String (Object)", "Ordinal", "0 (0.0%)", 5, "{Junior, Mid-Level, Senior, Lead/Principal, Executive}", "Organizational seniority hierarchy and professional responsibility tier.", "Mid-Level, Senior, Lead/Principal", "MVAR-Net: Career View / AAIE Decision Strategy"),
        ("Alumni_Profile.csv", "skills", "String (Object)", "Multi-Valued Nominal", "0 (0.0%)", 84, "Subset of 10 skills pool delimited by ';'", "Curated technical, leadership, and operational competencies claimed by alumni.", "Python;Machine Learning;Cloud Computing", "MVAR-Net: Career/Skill View (Multi-Hot Encoded)"),
        ("Alumni_Profile.csv", "base_affinity", "Float", "Continuous Ratio", "0 (0.0%)", 6000, "[0.080, 0.950]", "Latent intrinsic baseline affinity governing organic probability of interaction.", "0.8387, 0.1063, 0.8275", "Synthetic Generator Latent Parameter"),
        
        # Alumni_Interaction.csv
        ("Alumni_Interaction.csv", "alumni_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 5944, "Foreign key matching Alumni_Profile.csv", "Identifies the alumnus executing the respective platform action.", "ALU_00001, ALU_00003", "Relational Foreign Key (Temporal Tracking)"),
        ("Alumni_Interaction.csv", "interaction_month", "Integer", "Discrete Time", "0 (0.0%)", 24, "[1, 24] Calendar Months", "Sequential observation month timestamp (1-18: Obs Phase, 19-24: Evaluation Holdout).", "1, 11, 18, 22", "Temporal Window Filtering & Leakage Guard"),
        ("Alumni_Interaction.csv", "interaction_type", "String (Object)", "Nominal", "0 (0.0%)", 7, "{Login, Event_View, Event_Register, Event_Attend, Mentor_Request, Message_Sent, Networking_Connect}", "Modal event type defining the specific institutional engagement action.", "Login, Event_Attend, Mentor_Request", "EEN Behavioral Sequence Modeling / Feature Aggregation"),
        ("Alumni_Interaction.csv", "interaction_score", "Float", "Ratio", "0 (0.0%)", 6, "{1.0, 1.5, 2.0, 3.0, 4.0, 5.0}", "Assigned quantitative impact weight reflecting engagement intensity.", "1.0, 1.5, 3.0, 5.0", "MVAR-Net Behavioral View: Cumulative Score Sum"),
        
        # Event.csv
        ("Event.csv", "event_id", "String (Object)", "Nominal (Identifier)", "0 (0.0%)", 150, "EVT_0001 to EVT_0150", "Unique identifier for each institutional or alumni community event.", "EVT_0001, EVT_0002", "Primary Entity Key (Event Candidate Pool)"),
        ("Event.csv", "event_type", "String (Object)", "Nominal", "0 (0.0%)", 5, "{Webinar, Technical Workshop, Alumni Reunion, Career Fair, Leadership Panel}", "Pedagogical / social format of the scheduled event.", "Webinar, Technical Workshop", "RAOD Candidate Generation / ACRN Contextual Ranker"),
        ("Event.csv", "domain", "String (Object)", "Nominal", "0 (0.0%)", 6, "{Technology, Finance, Healthcare, Automotive, Consulting, Education}", "Core industry or disciplinary subject domain addressed by the event.", "Technology, Consulting, Finance", "RAOD Domain Filtering / ACRN Content Match"),
        ("Event.csv", "required_skills", "String (Object)", "Multi-Valued Nominal", "0 (0.0%)", 28, "2 skills from pool delimited by ';'", "Prerequisite or featured technical skills pertinent to the event curriculum.", "Cloud Computing;Python", "ACRN Skill-Compatibility Match Score"),
        ("Event.csv", "event_month", "Integer", "Discrete Time", "0 (0.0%)", 24, "[1, 24] Calendar Months", "Scheduled calendar month in which the event is hosted.", "5, 12, 19", "Temporal Recency Matching / Availability Filter"),
        ("Event.csv", "capacity", "Integer", "Discrete Ratio", "0 (0.0%)", 4, "{50, 100, 250, 500} Attendees", "Maximum allowable participant registration quota.", "50, 100, 250, 500", "AAIE Operational Constraint / Queue Management"),
        
        # Mentor.csv
        ("Mentor.csv", "mentor_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 600, "Matches senior alumni in Alumni_Profile", "Identifies the certified alumni mentor participating in the guidance pool.", "ALU_03506, ALU_05987", "Primary Key (Mentor Pool) / Foreign Key"),
        ("Mentor.csv", "industry", "String (Object)", "Nominal", "0 (0.0%)", 6, "{Technology, Finance, Healthcare, Automotive, Consulting, Education}", "Professional industry domain where the mentor exercises leadership.", "Technology, Consulting, Finance", "RAOD Mentor Candidate Match / ACRN Ranking"),
        ("Mentor.csv", "skills", "String (Object)", "Multi-Valued Nominal", "0 (0.0%)", 60, "Set of 3 expert skills delimited by ';'", "Domain expertise and leadership capabilities offered for mentee coaching.", "Leadership;Strategic Planning;Project Management", "ACRN Mentorship Skill-Transfer Scoring"),
        ("Mentor.csv", "seniority", "String (Object)", "Ordinal", "0 (0.0%)", 3, "{Senior, Lead/Principal, Executive}", "Professional seniority tier validating mentoring qualifications.", "Senior, Lead/Principal, Executive", "Eligibility Verification & Authority Weight"),
        ("Mentor.csv", "max_mentees", "Integer", "Discrete Ratio", "0 (0.0%)", 4, "{2, 3, 5, 8} Mentees", "Maximum number of simultaneous active mentees accepted by mentor.", "2, 3, 5, 8", "AAIE Mentorship Capacity Constraint"),
        ("Mentor.csv", "mentor_rating", "Float", "Interval", "0 (0.0%)", 71, "[4.30, 5.00] Stars", "Historical mentee satisfaction and peer evaluation score.", "4.45, 4.78, 5.00", "ACRN Quality Weight / Baseline Prioritization"),
        
        # Alumni_Network.csv
        ("Alumni_Network.csv", "source_alumni_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 5959, "Matches Alumni_Profile.csv", "Originating node in the directed peer networking relationship graph.", "ALU_00257, ALU_01916", "Relational Source Node (Graph Edge)"),
        ("Alumni_Network.csv", "target_alumni_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 5944, "Matches Alumni_Profile.csv", "Recipient node in the directed peer networking relationship graph.", "ALU_01992, ALU_03803", "Relational Target Node (Graph Edge)"),
        ("Alumni_Network.csv", "relationship_strength", "Float", "Continuous Ratio", "0 (0.0%)", 8, "[0.250, 1.000] ({0.25, 0.40, 0.50, 0.60, 0.65, 0.75, 0.85, 1.0})", "Calculated relationship strength based on shared major, industry, and archetype.", "0.25, 0.50, 0.85, 1.00", "DRMN: Edge Weight / Relational Feature View")
    ]
    
    cur_row = 5
    for row_data in master_features:
        ws2.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws2.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx in [1, 2] else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [5, 6]:
                cell.alignment = align_center
            elif col_idx in [3, 4]:
                cell.alignment = align_center
            elif col_idx in [1, 2]:
                cell.alignment = align_left
            else:
                cell.alignment = align_left
        cur_row += 1

    auto_fit_columns(ws2, min_widths={"A": 22, "B": 24, "C": 18, "D": 20, "E": 15, "F": 15, "G": 35, "H": 45, "I": 25, "J": 38}, max_width=48)

    # =========================================================================
    # SHEET 3: Alumni_Profile
    # =========================================================================
    ws3 = wb.create_sheet(title="Alumni_Profile")
    ws3.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws3,
        "ALUMNI PROFILE DATASET (Alumni_Profile.csv) - SCHEMA & EMPIRICAL DISTRIBUTIONS",
        "Demographic, Educational, Professional, and Competency Attributes (N = 6,000 Alumni Records)",
        9
    )
    
    headers3 = [
        "Attribute Name", "Data Type", "Scale", "Nulls (%)", "Distinct Values", 
        "Min / Max Domain", "Empirical Mean (Std) / Mode", "Operational Definition & Semantics", "Feature View Encoding"
    ]
    style_headers(ws3, 4, headers3)
    
    profile_details = [
        ("alumni_id", "String (Object)", "Nominal", "0 (0.0%)", 6000, "ALU_00001 to ALU_06000", "Unique Serial IDs", "Primary unique key identifying individual graduates in the institutional ecosystem.", "Entity Primary Key"),
        ("archetype", "Integer", "Nominal", "0 (0.0%)", 4, "0 to 3", "Mode: Archetype 0 (30.0%)", "Synthetic generative persona mapping career stage, skill portfolio, and engagement tendency.", "Ground-Truth Segmentation Reference"),
        ("graduation_year", "Integer", "Interval", "0 (0.0%)", 19, "2005 to 2023", "Mean: 2015.91 (Std: 5.12)", "Calendar year marking formal completion of academic program at the institution.", "MVAR-Net Profile View (Min-Max Scaled)"),
        ("degree", "String (Object)", "Nominal", "0 (0.0%)", 6, "6 Categories", "Mode: B.Tech (2,242; 37.4%)", "Academic qualification earned: B.Tech, M.Tech, MBA, B.Sc, M.Sc, Ph.D.", "MVAR-Net Profile View (One-Hot Encoded)"),
        ("major", "String (Object)", "Nominal", "0 (0.0%)", 6, "6 Disciplines", "Mode: Computer Science (2,452; 40.9%)", "Major academic specialization: CS, Data Science, Electrical Eng, Mech Eng, Business Admin, Bioeng.", "MVAR-Net Profile View / DRMN Relational Match"),
        ("industry", "String (Object)", "Nominal", "0 (0.0%)", 6, "6 Sectors", "Mode: Technology (3,022; 50.4%)", "Current professional vertical: Technology, Finance, Consulting, Automotive, Healthcare, Education.", "MVAR-Net Career View (One-Hot Encoded)"),
        ("years_experience", "Integer", "Ratio", "0 (0.0%)", 19, "1 to 19 Years", "Mean: 8.25 Years (Std: 4.96)", "Cumulative post-graduation career duration in professional workforce.", "MVAR-Net Career View (Min-Max Scaled)"),
        ("seniority", "String (Object)", "Ordinal", "0 (0.0%)", 5, "5 Tiers", "Mode: Mid-Level (2,150; 35.8%)", "Professional hierarchy status: Junior, Mid-Level, Senior, Lead/Principal, Executive.", "MVAR-Net Career View (One-Hot Encoded)"),
        ("skills", "String (Object)", "Multi-Valued", "0 (0.0%)", 84, "84 Combinations", "Mode: Project Mgmt; Finance Mod (117)", "Compound string of 2 to 3 claimed skills selected from the institutional skill pool.", "MVAR-Net Skill View (10 Multi-Hot Flags)"),
        ("base_affinity", "Float", "Ratio", "0 (0.0%)", 6000, "0.0800 to 0.9497", "Mean: 0.63 (Std: 0.24)", "Latent behavioral parameter governing interaction frequency in synthetic data generation.", "Synthetic Simulation Control Parameter")
    ]
    
    cur_row = 5
    for row_data in profile_details:
        ws3.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws3.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [2, 3, 4, 5]:
                cell.alignment = align_center
            elif col_idx == 1:
                cell.alignment = align_left
            else:
                cell.alignment = align_left
        cur_row += 1
        
    # Archetype Breakdown Sub-table
    cur_row += 2
    ws3.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=9)
    sec_cell = ws3.cell(row=cur_row, column=1, value="ALUMNI POPULATION ARCHETYPES (SYNTHETIC GENERATIVE SPECIFICATION)")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws3.row_dimensions[cur_row].height = 24
    
    cur_row += 1
    archetype_headers = ["Archetype ID", "Persona Designation", "Population Count (%)", "Graduation Window", "Experience Span", "Predominant Majors", "Predominant Industries", "Base Affinity Range", "Behavioral Tendency"]
    style_headers(ws3, cur_row, archetype_headers, fill=header_sub_fill)
    
    archetypes_data = [
        ("0", "Early-Career Tech Explorers", "1,800 (30.0%)", "2020 - 2023", "1 - 4 Years", "Computer Science, Data Science", "Technology (75%), Finance (25%)", "0.60 - 0.88", "High digital engagement, frequent event views, active workshop participation, seeking mentorship."),
        ("1", "Mid-Career Technical Specialists", "1,680 (28.0%)", "2014 - 2019", "5 - 10 Years", "CS, Data Science, Electrical Eng", "Technology (60%), Consulting (20%), Finance (20%)", "0.55 - 0.85", "Stable engagement, technical workshop attendance, targeted peer networking, potential mentors."),
        ("2", "Senior Mentors & Executive Leaders", "1,320 (22.0%)", "2005 - 2013", "11 - 19 Years", "CS, Business Admin, Electrical Eng", "Technology (50%), Consulting (30%), Finance (20%)", "0.68 - 0.95", "High institutional affinity, willing mentor pool, leadership panels, strong network centrality."),
        ("3", "At-Risk / Dormant Alumni", "1,200 (20.0%)", "2008 - 2020", "4 - 16 Years", "Mechanical Eng, Bioeng, Business Admin", "Automotive (40%), Healthcare (30%), Education (30%)", "0.08 - 0.30", "Very low engagement, sporadic logins, high disengagement risk (>90%), priority for AAIE re-engagement.")
    ]
    
    for row_data in archetypes_data:
        cur_row += 1
        ws3.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws3.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx in [1, 2] else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [1, 3]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    auto_fit_columns(ws3, min_widths={"A": 20, "B": 16, "C": 14, "D": 14, "E": 16, "F": 22, "G": 28, "H": 40, "I": 35}, max_width=45)

    # =========================================================================
    # SHEET 4: Alumni_Interaction
    # =========================================================================
    ws4 = wb.create_sheet(title="Alumni_Interaction")
    ws4.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws4,
        "ALUMNI INTERACTION LOGS (Alumni_Interaction.csv) - TEMPORAL ACTIVITY ARCHITECTURE",
        "Longitudinal Transactional Event Records (N = 107,487 Records Across 24 Calendar Months)",
        8
    )
    
    headers4 = [
        "Attribute Name", "Physical Type", "Measurement Scale", "Null Count (%)", 
        "Distinct Values", "Value Range / Categories", "Operational Definition & Semantics", "Downstream Pipeline Transformation"
    ]
    style_headers(ws4, 4, headers4)
    
    interaction_cols = [
        ("alumni_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 5944, "ALU_00001 to ALU_06000", "Identifies the specific alumnus performing the institutional event action.", "Group-by entity key for extracting 18-month longitudinal behavioral profile."),
        ("interaction_month", "Integer", "Discrete Time", "0 (0.0%)", 24, "Months 1 to 24", "Sequential monthly timestamp indicating the timing of the logged action.", "Chronological boundary split: Months 1-18 (Observation) vs. 19-24 (Holdout Test)."),
        ("interaction_type", "String (Object)", "Nominal", "0 (0.0%)", 7, "7 Specific Interaction Types", "Categorical identifier representing the modality of the alumni platform action.", "One-hot / count aggregation: event attendance, mentor requests, messaging counts."),
        ("interaction_score", "Float", "Ratio", "0 (0.0%)", 6, "{1.0, 1.5, 2.0, 3.0, 4.0, 5.0}", "Numerical engagement weight assigned to the interaction based on effort level.", "MVAR-Net Behavioral View: Aggregated sum_interaction_score feature.")
    ]
    
    cur_row = 5
    for row_data in interaction_cols:
        ws4.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws4.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [2, 3, 4, 5]:
                cell.alignment = align_center
            elif col_idx == 1:
                cell.alignment = align_left
            else:
                cell.alignment = align_left
        cur_row += 1
        
    # Interaction Modality Breakdown Table
    cur_row += 2
    ws4.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=8)
    sec_cell = ws4.cell(row=cur_row, column=1, value="INTERACTION MODALITY HIERARCHY & WEIGHTING SPECIFICATION")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws4.row_dimensions[cur_row].height = 24
    
    cur_row += 1
    modality_headers = ["Interaction Modality", "Base Weight Score", "Occurrence Count", "Percentage of Total (%)", "Engagement Effort Level", "Behavioral Interpretation in IAIF Framework"]
    style_headers(ws4, cur_row, modality_headers, fill=header_sub_fill)
    ws4.merge_cells(start_row=cur_row, start_column=6, end_row=cur_row, end_column=8)
    
    modalities = [
        ("Login", 1.0, 37440, "34.83%", "Passive / Baseline", "Basic portal authentication; reflects platform awareness and baseline activity."),
        ("Event_View", 1.5, 21352, "19.87%", "Mild Interest", "Browsing event schedules; signals thematic interest and discovery intent."),
        ("Event_Register", 3.0, 16244, "15.11%", "Active Commitment", "Formal registration for upcoming events; strong positive engagement signal."),
        ("Event_Attend", 5.0, 10885, "10.13%", "Deep Participation", "Actual validated attendance at events; highest behavioral commitment tier."),
        ("Message_Sent", 2.0, 10794, "10.04%", "Active Communication", "Peer-to-peer or institutional direct messaging; demonstrates social connectivity."),
        ("Networking_Connect", 3.0, 5399, "5.02%", "Social Expansion", "Initiating peer connection in alumni directory; expands professional network."),
        ("Mentor_Request", 4.0, 5373, "5.00%", "High-Value Mentorship", "Submitting structured mentorship application; high career growth motivation.")
    ]
    
    for row_data in modalities:
        cur_row += 1
        ws4.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        
        for col_idx in range(1, 6):
            cell = ws4.cell(row=cur_row, column=col_idx, value=row_data[col_idx - 1])
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [2, 4]:
                cell.alignment = align_center
            elif col_idx == 3:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            else:
                cell.alignment = align_left
                
        ws4.merge_cells(start_row=cur_row, start_column=6, end_row=cur_row, end_column=8)
        desc_cell = ws4.cell(row=cur_row, column=6, value=row_data[5])
        desc_cell.font = data_font
        desc_cell.fill = fill
        desc_cell.alignment = align_left
        for c_idx in range(6, 9):
            ws4.cell(row=cur_row, column=c_idx).border = cell_border
            ws4.cell(row=cur_row, column=c_idx).fill = fill
            
    # Temporal Split Breakdown Sub-table
    cur_row += 2
    ws4.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=8)
    sec_cell = ws4.cell(row=cur_row, column=1, value="LEAKAGE-FREE TEMPORAL PARTITIONING ARCHITECTURE")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws4.row_dimensions[cur_row].height = 24
    
    cur_row += 1
    temp_headers = ["Temporal Partition", "Calendar Horizon", "Interaction Records", "Percentage (%)", "Methodological Purpose & Leakage Guard"]
    style_headers(ws4, cur_row, temp_headers, fill=header_sub_fill)
    ws4.merge_cells(start_row=cur_row, start_column=5, end_row=cur_row, end_column=8)
    
    temp_data = [
        ("Observation Window (Input Space)", "Months 1 to 18", 83038, "77.25%", "All behavioral aggregates, interaction frequencies, and network metrics are computed strictly within this period to prevent temporal leakage into forecasting models."),
        ("Forecast Window (Target Space)", "Months 19 to 24", 24449, "22.75%", "Holdout 6-month evaluation horizon used exclusively to construct ground-truth disengagement labels (target_future_interactions == 0) and evaluate recommendation relevance.")
    ]
    
    for row_data in temp_data:
        cur_row += 1
        ws4.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        
        for col_idx in range(1, 5):
            cell = ws4.cell(row=cur_row, column=col_idx, value=row_data[col_idx - 1])
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [2, 4]:
                cell.alignment = align_center
            elif col_idx == 3:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            else:
                cell.alignment = align_left
                
        ws4.merge_cells(start_row=cur_row, start_column=5, end_row=cur_row, end_column=8)
        desc_cell = ws4.cell(row=cur_row, column=5, value=row_data[4])
        desc_cell.font = data_font
        desc_cell.fill = fill
        desc_cell.alignment = align_left
        for c_idx in range(5, 9):
            ws4.cell(row=cur_row, column=c_idx).border = cell_border
            ws4.cell(row=cur_row, column=c_idx).fill = fill

    auto_fit_columns(ws4, min_widths={"A": 22, "B": 18, "C": 18, "D": 15, "E": 18, "F": 22, "G": 40, "H": 40}, max_width=45)

    # =========================================================================
    # SHEET 5: Event
    # =========================================================================
    ws5 = wb.create_sheet(title="Event")
    ws5.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws5,
        "INSTITUTIONAL EVENTS DATASET (Event.csv) - SPECIFICATION & CANDIDATE CATALOG",
        "Curated Repository of Technical Workshops, Panels, and Reunions (N = 150 Events)",
        8
    )
    
    headers5 = [
        "Attribute Name", "Physical Type", "Measurement Scale", "Null Count (%)", 
        "Distinct Values", "Value Range / Categories", "Operational Definition & Semantics", "Recommender Pipeline Integration (RAOD/ACRN)"
    ]
    style_headers(ws5, 4, headers5)
    
    event_cols = [
        ("event_id", "String (Object)", "Nominal (Primary Key)", "0 (0.0%)", 150, "EVT_0001 to EVT_0150", "Unique system identifier for each institutional or alumni gathering.", "Target Item Identifier for Event Recommendation Task."),
        ("event_type", "String (Object)", "Nominal", "0 (0.0%)", 5, "{Webinar, Technical Workshop, Alumni Reunion, Career Fair, Leadership Panel}", "Format and organizational structure of the hosted event.", "Candidate filtering and user preference alignment matching."),
        ("domain", "String (Object)", "Nominal", "0 (0.0%)", 6, "{Technology, Finance, Healthcare, Automotive, Consulting, Education}", "Professional industry or disciplinary theme of the event.", "Domain-affinity compatibility scoring with alumni career industry."),
        ("required_skills", "String (Object)", "Multi-Valued", "0 (0.0%)", 28, "28 Combinations (2 skills each)", "Pair of focal competencies emphasized in the event curriculum.", "Content-based skill intersection matching with alumni profile skills."),
        ("event_month", "Integer", "Discrete Time", "0 (0.0%)", 24, "Months 1 to 24", "Scheduled calendar month for event execution.", "Temporal alignment with active alumni observation and holdout evaluation."),
        ("capacity", "Integer", "Discrete Ratio", "0 (0.0%)", 4, "{50, 100, 250, 500} Attendees", "Maximum attendee seating / registration limit.", "AAIE Decision Engine: Operational capacity constraint enforcement.")
    ]
    
    cur_row = 5
    for row_data in event_cols:
        ws5.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws5.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [2, 3, 4, 5]:
                cell.alignment = align_center
            elif col_idx == 1:
                cell.alignment = align_left
            else:
                cell.alignment = align_left
        cur_row += 1
        
    # Event Format Distribution Sub-table
    cur_row += 2
    ws5.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=8)
    sec_cell = ws5.cell(row=cur_row, column=1, value="EVENT FORMAT & INDUSTRIAL DOMAIN DISTRIBUTIONS")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws5.row_dimensions[cur_row].height = 24
    
    cur_row += 1
    style_headers(ws5, cur_row, ["Event Format", "Event Count", "Percentage (%)", "Primary Domain Focus", "Typical Capacity", "Target Audience Persona"], fill=header_sub_fill)
    ws5.merge_cells(start_row=cur_row, start_column=6, end_row=cur_row, end_column=8)
    
    event_breakdown = [
        ("Webinar", 41, "27.33%", "Technology, Finance, Consulting", "100 - 500", "Broad alumni base, remote attendees, exploratory learners."),
        ("Technical Workshop", 34, "22.67%", "Technology (Python, ML, Cloud)", "50 - 100", "Early to mid-career engineers seeking concrete upskilling."),
        ("Alumni Reunion", 33, "22.00%", "Multi-Domain Institutional", "250 - 500", "Cross-cohort networking, institutional milestone cohorts."),
        ("Career Fair", 21, "14.00%", "Technology, Finance, Consulting", "250 - 500", "Recent graduates, junior alumni seeking lateral transitions."),
        ("Leadership Panel", 21, "14.00%", "Consulting, Finance, Technology", "50 - 100", "Senior executives, lead practitioners, strategic management.")
    ]
    
    for row_data in event_breakdown:
        cur_row += 1
        ws5.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx in range(1, 6):
            cell = ws5.cell(row=cur_row, column=col_idx, value=row_data[col_idx - 1])
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [3, 5]:
                cell.alignment = align_center
            elif col_idx == 2:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            else:
                cell.alignment = align_left
                
        ws5.merge_cells(start_row=cur_row, start_column=6, end_row=cur_row, end_column=8)
        desc_cell = ws5.cell(row=cur_row, column=6, value=row_data[5])
        desc_cell.font = data_font
        desc_cell.fill = fill
        desc_cell.alignment = align_left
        for c_idx in range(6, 9):
            ws5.cell(row=cur_row, column=c_idx).border = cell_border
            ws5.cell(row=cur_row, column=c_idx).fill = fill

    auto_fit_columns(ws5, min_widths={"A": 22, "B": 18, "C": 18, "D": 15, "E": 18, "F": 25, "G": 38, "H": 38}, max_width=45)

    # =========================================================================
    # SHEET 6: Mentor
    # =========================================================================
    ws6 = wb.create_sheet(title="Mentor")
    ws6.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws6,
        "ALUMNI MENTOR ROSTER (Mentor.csv) - SCHEMA & MENTORSHIP POOL CHARACTERISTICS",
        "Expert Mentor Profiles, Capacity Thresholds, and Satisfaction Ratings (N = 600 Mentors)",
        8
    )
    
    headers6 = [
        "Attribute Name", "Physical Type", "Measurement Scale", "Null Count (%)", 
        "Distinct Values", "Value Range / Categories", "Operational Definition & Semantics", "Recommender Pipeline Integration (RAOD/ACRN)"
    ]
    style_headers(ws6, 4, headers6)
    
    mentor_cols = [
        ("mentor_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 600, "Matches Senior Alumni_Profile IDs", "Primary mentor identifier referencing the senior alumni profile record.", "Target Item Identifier for Mentor Recommendation Task."),
        ("industry", "String (Object)", "Nominal", "0 (0.0%)", 6, "{Technology, Finance, Healthcare, Automotive, Consulting, Education}", "Professional industry vertical in which the mentor exercises leadership.", "Domain compatibility matching with junior/seeking alumni."),
        ("skills", "String (Object)", "Multi-Valued", "0 (0.0%)", 60, "60 Skill Triplets", "Specialized technical, operational, and strategic guidance domains.", "ACRN Mentorship Relevance Score: Skill transfer match."),
        ("seniority", "String (Object)", "Ordinal", "0 (0.0%)", 3, "{Senior, Lead/Principal, Executive}", "Professional seniority tier validating mentoring qualifications.", "Hierarchy constraint verification (mentor seniority > mentee seniority)."),
        ("max_mentees", "Integer", "Discrete Ratio", "0 (0.0%)", 4, "{2, 3, 5, 8} Mentees", "Maximum number of simultaneous active mentees accepted.", "AAIE Decision Engine: Capacity constraint dispatching."),
        ("mentor_rating", "Float", "Interval", "0 (0.0%)", 71, "[4.30, 5.00] Stars", "Historical mentee satisfaction rating score.", "ACRN Ranking Weight: Quality prior in mentor ranking.")
    ]
    
    cur_row = 5
    for row_data in mentor_cols:
        ws6.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws6.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [2, 3, 4, 5]:
                cell.alignment = align_center
            elif col_idx == 1:
                cell.alignment = align_left
            else:
                cell.alignment = align_left
        cur_row += 1
        
    # Mentor Seniority and Capacity Sub-table
    cur_row += 2
    ws6.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=8)
    sec_cell = ws6.cell(row=cur_row, column=1, value="MENTOR SENIORITY TIERS, CAPACITY LIMITS, AND SATISFACTION METRICS")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws6.row_dimensions[cur_row].height = 24
    
    cur_row += 1
    style_headers(ws6, cur_row, ["Seniority Tier", "Mentor Count", "Percentage (%)", "Mean Mentor Rating", "Standard Deviation", "Capacity Range (Mentees)", "Mentorship Focus Area"], fill=header_sub_fill)
    ws6.merge_cells(start_row=cur_row, start_column=7, end_row=cur_row, end_column=8)
    
    mentor_tiers = [
        ("Senior", 323, "53.83%", "4.65 Stars", "0.20", "2 - 5 Mentees", "Career transition, practical technical mentoring, project leadership."),
        ("Lead / Principal", 193, "32.17%", "4.66 Stars", "0.20", "3 - 8 Mentees", "Architectural guidance, technical strategy, cross-functional leadership."),
        ("Executive", 84, "14.00%", "4.66 Stars", "0.21", "2 - 5 Mentees", "C-suite coaching, entrepreneurship, strategic institutional philanthropy.")
    ]
    
    for row_data in mentor_tiers:
        cur_row += 1
        ws6.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx in range(1, 7):
            cell = ws6.cell(row=cur_row, column=col_idx, value=row_data[col_idx - 1])
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [3, 4, 5, 6]:
                cell.alignment = align_center
            elif col_idx == 2:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            else:
                cell.alignment = align_left
                
        ws6.merge_cells(start_row=cur_row, start_column=7, end_row=cur_row, end_column=8)
        desc_cell = ws6.cell(row=cur_row, column=7, value=row_data[6])
        desc_cell.font = data_font
        desc_cell.fill = fill
        desc_cell.alignment = align_left
        for c_idx in range(7, 9):
            ws6.cell(row=cur_row, column=c_idx).border = cell_border
            ws6.cell(row=cur_row, column=c_idx).fill = fill

    auto_fit_columns(ws6, min_widths={"A": 22, "B": 18, "C": 18, "D": 15, "E": 18, "F": 25, "G": 25, "H": 35}, max_width=45)

    # =========================================================================
    # SHEET 7: Alumni_Network
    # =========================================================================
    ws7 = wb.create_sheet(title="Alumni_Network")
    ws7.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws7,
        "ALUMNI NETWORK TOPOLOGY (Alumni_Network.csv) - RELATIONAL GRAPH SCHEMA",
        "Peer Connectivity Edges and Dependency-Aware Affinity Strengths (N = 29,984 Directed Edges)",
        7
    )
    
    headers7 = [
        "Attribute Name", "Physical Type", "Measurement Scale", "Null Count (%)", 
        "Distinct Values", "Value Range / Categories", "Operational Definition & Graph Modeling Semantics"
    ]
    style_headers(ws7, 4, headers7)
    
    network_cols = [
        ("source_alumni_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 5959, "Matches Alumni_Profile.csv", "Originating alumnus node initiating or anchoring the relational graph edge."),
        ("target_alumni_id", "String (Object)", "Nominal (Foreign Key)", "0 (0.0%)", 5944, "Matches Alumni_Profile.csv", "Recipient alumnus node connected in the institutional social network graph."),
        ("relationship_strength", "Float", "Continuous Ratio", "0 (0.0%)", 8, "[0.250, 1.000]", "Calculated edge affinity weight quantifying connection intimacy and mutual professional relevance.")
    ]
    
    cur_row = 5
    for row_data in network_cols:
        ws7.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx, val in enumerate(row_data, 1):
            cell = ws7.cell(row=cur_row, column=col_idx, value=val)
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [2, 3, 4, 5]:
                cell.alignment = align_center
            elif col_idx == 1:
                cell.alignment = align_left
            else:
                cell.alignment = align_left
        cur_row += 1
        
    # Edge Weight Formulation Rules Sub-table
    cur_row += 2
    ws7.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=7)
    sec_cell = ws7.cell(row=cur_row, column=1, value="RELATIONSHIP STRENGTH MATHEMATICAL FORMULATION & RULES")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws7.row_dimensions[cur_row].height = 24
    
    cur_row += 1
    rules_headers = ["Affinity Component", "Mathematical Contribution", "Condition / Criterion", "Contribution Rationale & Graph Modeling Impact"]
    style_headers(ws7, cur_row, rules_headers, fill=header_sub_fill)
    ws7.merge_cells(start_row=cur_row, start_column=4, end_row=cur_row, end_column=7)
    
    rules = [
        ("Base Institutional Connection", "+0.25", "Default for all connected edge pairs", "Represents foundational mutual affiliation through graduation from the same parent university."),
        ("Shared Academic Major Bonus", "+0.35", "source.major == target.major", "Strong disciplinary affinity: shared curriculum, common faculty, and similar foundational technical background."),
        ("Shared Industry Sector Bonus", "+0.25", "source.industry == target.industry", "Professional relevance bonus: identical employment domain, shared market ecosystem, and lateral networking value."),
        ("Shared Persona Archetype Bonus", "+0.15", "source.archetype == target.archetype", "Generational and behavioral alignment: similar career stage, seniority expectations, and platform participation habit."),
        ("Edge Weight Normalization", "min(1.0, sum)", "Bounded Upper Limit", "Caps maximum relationship strength at 1.000 for strict numerical stability in graph neural architectures.")
    ]
    
    for row_data in rules:
        cur_row += 1
        ws7.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx in range(1, 4):
            cell = ws7.cell(row=cur_row, column=col_idx, value=row_data[col_idx - 1])
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx == 2:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
        ws7.merge_cells(start_row=cur_row, start_column=4, end_row=cur_row, end_column=7)
        desc_cell = ws7.cell(row=cur_row, column=4, value=row_data[3])
        desc_cell.font = data_font
        desc_cell.fill = fill
        desc_cell.alignment = align_left
        for c_idx in range(4, 8):
            ws7.cell(row=cur_row, column=c_idx).border = cell_border
            ws7.cell(row=cur_row, column=c_idx).fill = fill
            
    # Empirical Distribution of Weights
    cur_row += 2
    ws7.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=7)
    sec_cell = ws7.cell(row=cur_row, column=1, value="EMPIRICAL DISTRIBUTION OF RELATIONSHIP STRENGTHS")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws7.row_dimensions[cur_row].height = 24
    
    cur_row += 1
    style_headers(ws7, cur_row, ["Edge Strength Value", "Edge Count", "Percentage (%)", "Cumulative (%)", "Structural Meaning in Graph"], fill=header_sub_fill)
    ws7.merge_cells(start_row=cur_row, start_column=5, end_row=cur_row, end_column=7)
    
    weight_dist = [
        ("0.250", 14151, "47.20%", "47.20%", "General campus acquaintance (No shared major, industry, or archetype)."),
        ("0.400", 2403, "8.01%", "55.21%", "Shared archetype only (+0.15)."),
        ("0.500", 3842, "12.81%", "68.02%", "Shared industry only (+0.25)."),
        ("0.600", 2443, "8.15%", "76.17%", "Shared major only (+0.35)."),
        ("0.650", 2104, "7.02%", "83.19%", "Shared industry and shared archetype (+0.40)."),
        ("0.750", 1642, "5.48%", "88.67%", "Shared major and shared archetype (+0.50)."),
        ("0.850", 1780, "5.94%", "94.61%", "Shared major and shared industry (+0.60)."),
        ("1.000", 1619, "5.40%", "100.00%", "Perfect alignment: Shared major, shared industry, and shared archetype.")
    ]
    
    for row_data in weight_dist:
        cur_row += 1
        ws7.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        for col_idx in range(1, 5):
            cell = ws7.cell(row=cur_row, column=col_idx, value=row_data[col_idx - 1])
            cell.font = data_bold_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [1, 3, 4]:
                cell.alignment = align_center
            elif col_idx == 2:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            else:
                cell.alignment = align_left
                
        ws7.merge_cells(start_row=cur_row, start_column=5, end_row=cur_row, end_column=7)
        desc_cell = ws7.cell(row=cur_row, column=5, value=row_data[4])
        desc_cell.font = data_font
        desc_cell.fill = fill
        desc_cell.alignment = align_left
        for c_idx in range(5, 8):
            ws7.cell(row=cur_row, column=c_idx).border = cell_border
            ws7.cell(row=cur_row, column=c_idx).fill = fill

    auto_fit_columns(ws7, min_widths={"A": 22, "B": 18, "C": 18, "D": 15, "E": 18, "F": 25, "G": 45}, max_width=45)

    # =========================================================================
    # SHEET 8: Engineered_Features_IAIF
    # =========================================================================
    ws8 = wb.create_sheet(title="Engineered_Features_IAIF")
    ws8.views.sheetView[0].showGridLines = True
    
    apply_title_banner(
        ws8,
        "IAIF ENGINEERED MULTI-VIEW FEATURE SPACES & SUPERVISED TARGETS",
        "Fused Behavioral Vectors, Multi-Modal Encodings, and Algorithmic Target Definitions",
        8
    )
    
    headers8 = [
        "Feature Space / Vector", "Feature Name", "Dimensionality / Type", "Extraction Formula / Source", 
        "Normalization / Encoding", "Core Algorithmic Role in IAIF Architecture"
    ]
    style_headers(ws8, 4, headers8)
    
    engineered = [
        # Profile View
        ("Profile View (dim=14)", "years_experience_norm", "1 Dim Continuous", "Alumni_Profile.csv", "MinMax(years_experience, [0, 1])", "Demographic career maturity representation in MVAR-Net."),
        ("Profile View (dim=14)", "graduation_year_norm", "1 Dim Continuous", "Alumni_Profile.csv", "MinMax(graduation_year, [0, 1])", "Temporal cohort anchoring in MVAR-Net Profile Encoder."),
        ("Profile View (dim=14)", "degree_ohe_[6]", "6 Dim Binary", "Alumni_Profile.csv", "One-Hot Encoding {BTech, MTech, MBA, BSc, MSc, PhD}", "Academic credential categorical encoding."),
        ("Profile View (dim=14)", "major_ohe_[6]", "6 Dim Binary", "Alumni_Profile.csv", "One-Hot Encoding {CS, DS, EE, ME, BA, BioE}", "Disciplinary specialization representation."),
        
        # Career & Skill View
        ("Career & Skill View (dim=22)", "years_experience_norm", "1 Dim Continuous", "Alumni_Profile.csv", "MinMax(years_experience, [0, 1])", "Workplace experience duration."),
        ("Career & Skill View (dim=22)", "industry_ohe_[6]", "6 Dim Binary", "Alumni_Profile.csv", "One-Hot Encoding (6 Industries)", "Employment domain representation."),
        ("Career & Skill View (dim=22)", "seniority_ohe_[5]", "5 Dim Binary", "Alumni_Profile.csv", "One-Hot Encoding (5 Seniority Tiers)", "Organizational responsibility tier."),
        ("Career & Skill View (dim=22)", "skills_multi_hot_[10]", "10 Dim Binary", "Alumni_Profile.csv", "Multi-Hot Indicator across SKILLS_POOL", "Technical and strategic competency vector."),
        
        # Behavioral View
        ("Behavioral & Relational View (dim=9)", "total_interactions", "1 Dim Ratio", "Count(actions in months 1-18)", "MinMax Scaled [0, 1]", "Overall platform activity volume in observation period."),
        ("Behavioral & Relational View (dim=9)", "sum_interaction_score", "1 Dim Ratio", "Sum(scores in months 1-18)", "MinMax Scaled [0, 1]", "Effort-weighted interaction engagement volume."),
        ("Behavioral & Relational View (dim=9)", "recency_months", "1 Dim Interval", "18 - max(interaction_month)", "MinMax Scaled [0, 1]", "Inactivity duration (recency of last observed interaction)."),
        ("Behavioral & Relational View (dim=9)", "unique_types", "1 Dim Discrete", "Unique(interaction_types)", "MinMax Scaled [0, 1]", "Behavioral diversity across platforms features."),
        ("Behavioral & Relational View (dim=9)", "event_attendance", "1 Dim Discrete", "Count(Event_Attend)", "MinMax Scaled [0, 1]", "Direct participation in organized institutional events."),
        ("Behavioral & Relational View (dim=9)", "mentor_requests", "1 Dim Discrete", "Count(Mentor_Request)", "MinMax Scaled [0, 1]", "Active interest in receiving structured mentorship."),
        ("Behavioral & Relational View (dim=9)", "messages_sent", "1 Dim Discrete", "Count(Message_Sent)", "MinMax Scaled [0, 1]", "Peer-to-peer communicative activity."),
        ("Behavioral & Relational View (dim=9)", "network_degree", "1 Dim Discrete", "Degree in Alumni_Network.csv", "MinMax Scaled [0, 1]", "Alumni social network centrality and node degree."),
        ("Behavioral & Relational View (dim=9)", "mean_rel_strength", "1 Dim Ratio", "Mean(edge weights)", "MinMax Scaled [0, 1]", "Average relational affinity to immediate peer connections."),
        
        # Central Unified Latent Representation
        ("Central Unified Latent Space", "Latent Vector z", "64 Dim Continuous", "MVAR-Net Attention-Weighted Multi-View Fusion", "LayerNorm + Mish Activation", "Fused representation driving BAAC Clustering, EEN Prediction, and Recommenders."),
        
        # Algorithmic Prediction Targets
        ("Prediction Target (EEN Binary)", "disengagement_target", "1 Dim Binary {0, 1}", "1 if interactions in months 19-24 == 0 else 0", "Ground-Truth Class Label", "Primary supervised label for 6-month holdout disengagement forecasting."),
        ("Prediction Target (EEN Multi-Class)", "engagement_state", "1 Dim Ordinal {0, 1, 2}", "Tier: 2 if score>=30, 1 if score>=10, else 0", "Ground-Truth Class Label", "Supervised target for multi-tier engagement state classification.")
    ]
    
    cur_row = 5
    for row_data in engineered:
        ws8.row_dimensions[cur_row].height = 24
        fill = alt_fill if cur_row % 2 == 0 else white_fill
        
        for col_idx in range(1, 6):
            cell = ws8.cell(row=cur_row, column=col_idx, value=row_data[col_idx - 1])
            cell.font = data_bold_font if col_idx in [1, 2] else data_font
            cell.fill = fill
            cell.border = cell_border
            if col_idx in [3]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
        ws8.merge_cells(start_row=cur_row, start_column=6, end_row=cur_row, end_column=8)
        desc_cell = ws8.cell(row=cur_row, column=6, value=row_data[5])
        desc_cell.font = data_font
        desc_cell.fill = fill
        desc_cell.alignment = align_left
        for c_idx in range(6, 9):
            ws8.cell(row=cur_row, column=c_idx).border = cell_border
            ws8.cell(row=cur_row, column=c_idx).fill = fill
            
        cur_row += 1

    auto_fit_columns(ws8, min_widths={"A": 28, "B": 24, "C": 18, "D": 32, "E": 32, "F": 45}, max_width=48)

    # -------------------------------------------------------------
    # Save Workbook to both root and results/tables
    # -------------------------------------------------------------
    os.makedirs("results/tables", exist_ok=True)
    root_path = "Dataset_Description_Table.xlsx"
    tables_path = "results/tables/Dataset_Description_Table.xlsx"
    
    wb.save(root_path)
    wb.save(tables_path)
    print(f"Successfully generated and saved Excel workbook at:\n  1. {os.path.abspath(root_path)}\n  2. {os.path.abspath(tables_path)}")

if __name__ == "__main__":
    build_excel()
